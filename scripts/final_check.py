import openpyxl
wb = openpyxl.load_workbook('digital-products/hdb-renosmart-planner/HDB-RenoSmart-Budget-Compliance-Planner-2026.xlsx')

# Final comprehensive check
print('=== FINAL COMPREHENSIVE CHECK ===')
print()

# Check for any formula errors
errors = []
for name in wb.sheetnames:
    ws = wb[name]
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
        for c in row:
            if c.value and isinstance(c.value, str) and c.value.startswith('='):
                v = c.value
                if '#REF' in v or '#VALUE' in v or '#NAME' in v or '#DIV' in v or '#N/A' in v:
                    errors.append(name + '!' + c.coordinate + ': ' + v)
if errors:
    print('ERRORS FOUND:')
    for e in errors:
        print('  ' + e)
else:
    print('No formula errors')

print()
print('Sheets:', wb.sheetnames)
print()

# Check product config
with open('src/config/products.ts', 'r') as f:
    content = f.read()
    if 'HDB_RENOSMART_PRODUCT' in content:
        print('Product config: HDB_RENOSMART_PRODUCT present')
        if 'S$19' in content or 'priceAmount: 19' in content:
            print('  Price: S$19 confirmed')
        if 'hdb-renovation-cost-guide-singapore' in content:
            print('  Article slug: hdb-renovation-cost-guide-singapore confirmed')
        if 'PLACEHOLDER_HDB_RENOSMART' in content:
            print('  PayPal URL: PLACEHOLDER (not yet configured)')

# Check CTA
with open('src/components/PaidProductCTA.tsx', 'r') as f:
    content = f.read()
    if 'hdb-renosmart-planner' in content:
        print('CTA: hdb-renosmart-planner present')
        if 'Before you renovate, know the true cost.' in content:
            print('  CTA copy present')
        if 'Personalised renovation budget' in content:
            print('  Value bullets present')

print()
print('=== ALL CHECKS COMPLETE ===')